import io

import docx
import pymupdf
import pytest
from fastapi.testclient import TestClient

from app.main import app
from tests.factories import make_resume

client = TestClient(app)
resume_text = (
    "Asha Patel\nasha@example.com\nEXPERIENCE\nSoftware Engineer, Acme 2023 to Present\n"
    + "Built a FastAPI service handling 2000 requests per minute for internal tools.\n" * 6
    + "SKILLS\nPython, TypeScript, PostgreSQL\nEDUCATION\nB.E. IT, GTU"
)


def _pdf_bytes(text: str) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(40, 40, 560, 800), text, fontsize=9)
    page.insert_link(
        {
            "kind": pymupdf.LINK_URI,
            "from": pymupdf.Rect(40, 780, 100, 790),
            "uri": "https://github.com/asha",
        }
    )
    return doc.tobytes()


def _docx_bytes(text: str) -> bytes:
    document = docx.Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Certifications"
    table.cell(0, 1).text = "AWS Cloud Practitioner"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_health():
    assert client.get("/api/v1/health").json() == {"status": "ok"}


def test_meta_exposes_limits():
    body = client.get("/api/v1/meta").json()
    assert body["maxUploadBytes"] > 0 and ".pdf" in body["acceptedTypes"]


def test_parse_pdf_extracts_text_and_link_annotations():
    r = client.post(
        "/api/v1/resumes/parse",
        files={"file": ("cv.pdf", _pdf_bytes(resume_text), "application/pdf")},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["fileType"] == "pdf" and body["pageCount"] == 1
    assert "FastAPI" in body["text"]
    assert "https://github.com/asha" in body["text"]


def test_parse_docx_reads_tables():
    r = client.post("/api/v1/resumes/parse", files={"file": ("cv.docx", _docx_bytes(resume_text))})
    assert r.status_code == 200, r.text
    assert "AWS Cloud Practitioner" in r.json()["text"]


@pytest.mark.parametrize(
    ("name", "data", "code"),
    [
        ("cv.txt", b"hello", "unsupportedFileType"),
        ("cv.doc", b"hello", "unsupportedFileType"),
        ("cv.pdf", b"not a pdf", "corruptFile"),
        ("cv.docx", b"PK\x03\x04broken", "corruptFile"),
        ("cv.pdf", b"", "emptyFile"),
    ],
)
def test_parse_rejects_bad_files_with_clear_codes(name, data, code):
    r = client.post("/api/v1/resumes/parse", files={"file": (name, data)})
    assert r.status_code >= 400
    error = r.json()["error"]
    assert error["code"] == code and error["message"] and error["hint"]


def test_parse_rejects_image_only_pdf():
    doc = pymupdf.open()
    doc.new_page()
    r = client.post("/api/v1/resumes/parse", files={"file": ("cv.pdf", doc.tobytes())})
    assert r.json()["error"]["code"] == "noTextFound"


def test_parse_rejects_oversized_upload():
    big = b"%PDF-" + b"0" * (6 * 1024 * 1024)
    r = client.post("/api/v1/resumes/parse", files={"file": ("cv.pdf", big)})
    assert r.status_code == 413 and r.json()["error"]["code"] == "fileTooLarge"


def test_tailor_rejects_short_job_description_as_stream_error():
    r = client.post(
        "/api/v1/tailor", json={"jobDescription": "too short", "resumeText": resume_text}
    )
    assert "event: error" in r.text and "jobDescriptionTooShort" in r.text


@pytest.mark.parametrize("fmt", ["pdf", "docx", "txt"])
def test_export_formats(fmt):
    resume = make_resume().model_dump(by_alias=True)
    r = client.post("/api/v1/export", json={"format": fmt, "resume": resume})
    assert r.status_code == 200
    assert "Asha-Patel-Resume" in r.headers["content-disposition"]
    if fmt == "pdf":
        pdf = pymupdf.open(stream=r.content, filetype="pdf")
        text = pdf[0].get_text()
        assert "Asha Patel" in text and "EXPERIENCE" in text.upper()


def test_export_cover_letter_requires_text():
    resume = make_resume().model_dump(by_alias=True)
    r = client.post(
        "/api/v1/export", json={"kind": "coverLetter", "format": "pdf", "resume": resume}
    )
    assert r.json()["error"]["code"] == "emptyCoverLetter"
